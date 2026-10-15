from __future__ import annotations

import datetime as dt
import io
from types import SimpleNamespace

import pytest

from app.models.schema import Column
from app.models.spec import ProcessSpec
from app.services.schema_infer import (
    columns_from_ddl,
    columns_from_json_text,
    columns_from_layout,
    compare_with_sample,
    fixed_width_rows,
    guess_types_with_bigquery,
    infer_from_bytes,
    lab_guess_target,
    preview_fixed_width,
)
from app.services.validators import ClientFactory, SpecValidator, order_checks

from conftest import load_sample, minimal_spec
from fakes import FakeBq, FakeGcs

RUN_DATE = dt.date(2026, 1, 14)
CSV = b"store_id,sku,sold_at,sale_date,quantity,amount,store_code\n1,A1,2026-01-14 10:00:00+00:00,2026-01-14,3,9.99,ABC1\n"


def test_csv_inference_types_and_renames():
    r = infer_from_bytes(b"Store ID,zip,amount,ok,day\n1,01234,9.5,true,2026-01-01\n", "CSV")
    assert [(c.name, c.type) for c in r.columns] == [
        ("Store_ID", "INTEGER"), ("zip", "STRING"), ("amount", "FLOAT"), ("ok", "BOOLEAN"), ("day", "DATE"),
    ]
    assert r.warnings  # header renamed


def test_ndjson_nested_and_repeated():
    r = infer_from_bytes(b'{"id": 1, "tags": ["a"], "addr": {"zip": "02134"}}\n', "NEWLINE_DELIMITED_JSON")
    cols = {c.name: c for c in r.columns}
    assert cols["tags"].mode == "REPEATED" and cols["addr"].type == "RECORD"
    assert cols["addr"].fields[0].type == "STRING"  # JSON strings stay strings


def test_json_array_gets_a_helpful_error():
    with pytest.raises(ValueError, match="newline-delimited"):
        infer_from_bytes(b'[{"a": 1},\n{"a": 2}]', "NEWLINE_DELIMITED_JSON")


def test_ddl_and_json_schema_sources():
    cols = columns_from_ddl("CREATE TABLE t (id INT64 NOT NULL, tags ARRAY<STRING>, s STRUCT<a STRING>)")
    assert [(c.name, c.type, c.mode) for c in cols] == [("id", "INTEGER", "REQUIRED"), ("tags", "STRING", "REPEATED"), ("s", "RECORD", "NULLABLE")]
    assert columns_from_json_text('{"fields": [{"name": "a", "type": "INT64"}]}')[0].type == "INTEGER"


def test_compare_with_sample_flags_order_and_types():
    sample = infer_from_bytes(b"b,a\nx,1\n", "CSV")
    problems = compare_with_sample([Column(name="a", type="INTEGER"), Column(name="b", type="STRING")], sample, "CSV")
    assert any("order" in p for p in problems)


PROD_NAMES = ("acme-analytics-prod", "acme-landing-prod", "acme-exports-prod", "acme-crm-dropzone-prod")


def _validator(gcs=None, bq=None, **kw):
    return SpecValidator(ClientFactory(bq=bq or FakeBq(datasets={"acme-analytics-dev.raw_sales", "acme-analytics-dev.sales_mart"}),
                                       gcs=gcs or FakeGcs()), **kw)


def _touched(bq: FakeBq, gcs: FakeGcs) -> str:
    return " ".join([*bq.refs, *(sql for sql, _cfg in bq.queries), *gcs.calls])


def test_order_check_flags_reading_a_table_built_later(sales_spec_dict):
    sales_spec_dict["steps"] = [sales_spec_dict["steps"][i] for i in (0, 2, 1, 3)]  # export before the MERGE
    items = order_checks(ProcessSpec.model_validate(sales_spec_dict))
    warn = [i for i in items if i.level == "warn"]
    assert warn and warn[0].fix == {"action": "move_after", "step": "export_sales_daily_to_finance", "after": "build_sales_daily_agg"}


def test_order_check_ok_for_sample():
    items = order_checks(ProcessSpec.model_validate(load_sample("sales_daily")))
    assert [i.level for i in items] == ["ok"]


def test_full_validation_with_fake_gcp_uses_lab_values_only():
    gcs = FakeGcs({
        "acme-landing-dev": {"sales/incoming/store_sales_20260114_part1.csv": CSV, "sales/archive/store_sales_old.csv": CSV},
        "acme-exports-dev": {},
    })
    bq = FakeBq(datasets={"acme-analytics-dev.raw_sales"})
    report = _validator(gcs=gcs, bq=bq).validate(ProcessSpec.model_validate(load_sample("sales_daily")))
    titles = [i.title for i in report.items]
    assert report.environment == "lab" and titles[0].startswith("Checked in lab only")
    assert any("1 file(s) would be loaded in lab" in t for t in titles)  # the archived file is not counted
    assert any("Table schema matches the sample file" in t for t in titles)
    assert any("dry run OK in lab" in t for t in titles)
    assert any("does not exist yet in lab; the first run creates it" in t for t in titles)
    assert any("Code generates cleanly" in t for t in titles)
    assert report.ok, [i for i in report.items if i.level == "error"]
    touched = _touched(bq, gcs)
    assert "acme-analytics-dev" in touched and not any(p in touched for p in PROD_NAMES)


def test_missing_dataset_and_no_files_are_reported():
    bq = FakeBq(datasets=set())
    report = _validator(gcs=FakeGcs({"acme-landing-dev": {"sales/incoming/other.csv": CSV}}), bq=bq).validate(
        ProcessSpec.model_validate(load_sample("sales_daily"))
    )
    errors = [i.title for i in report.items if i.level == "error"]
    infos = [i.title for i in report.items if i.level == "info"]
    assert any("dataset acme-analytics-dev.raw_sales does not exist in lab" in e for e in errors)
    assert any("No file named '*store_sales_*.csv'" in t for t in infos)
    assert not report.ok


def test_target_table_columns_are_compared():
    table = SimpleNamespace(schema=[SimpleNamespace(name=n, field_type=t, mode="NULLABLE") for n, t in (
        ("store_id", "INTEGER"), ("sku", "STRING"), ("sold_at", "TIMESTAMP"), ("sale_date", "DATE"), ("quantity", "INTEGER"),
        ("amount", "STRING"), ("file_batch_id", "STRING"), ("inserted_ts", "DATETIME"), ("updated_ts", "DATETIME"),
        ("file_name", "STRING"), ("legacy", "STRING"))])
    bq = FakeBq(datasets={"acme-analytics-dev.raw_sales"}, tables={"acme-analytics-dev.raw_sales.store_sales": table})
    report = _validator(bq=bq).validate(ProcessSpec.model_validate(load_sample("sales_daily")))
    titles = " | ".join(i.title for i in report.items if i.step == "ingest_store_sales_files")
    assert "Columns not in the existing target table: store_code" in titles
    assert "columns the file does not: legacy" in titles and "Type differences" in titles


def test_dry_run_failure_is_an_error_but_tables_built_earlier_are_info():
    bq = FakeBq(datasets={"acme-analytics-dev.raw_sales"}, failing_sql={"MERGE": "400 Syntax error: Unexpected keyword"})
    report = _validator(bq=bq).validate(ProcessSpec.model_validate(load_sample("sales_daily")))
    assert any(i.level == "error" and "dry run failed" in i.title for i in report.items)

    bq2 = FakeBq(datasets={"acme-analytics-dev.crm"}, failing_sql={"dim_customer": "404 Not found: Table acme-analytics-dev:crm.customer_stg"})
    spec = ProcessSpec.model_validate(load_sample("customer_snapshot"))
    report2 = SpecValidator(ClientFactory(bq=bq2, gcs=FakeGcs())).validate(spec)
    merge_items = [i for i in report2.items if i.step == "merge_customer_dimension" and "dry run" in i.title]
    assert merge_items and merge_items[0].level == "info"


def test_csv_export_of_nested_query_is_an_error():
    spec = ProcessSpec.model_validate(load_sample("sales_daily"))
    bq = FakeBq(datasets={"acme-analytics-dev.raw_sales"}, schema=[SimpleNamespace(name="items", field_type="RECORD", mode="REPEATED")])
    report = _validator(bq=bq).validate(spec)
    assert any("CSV cannot hold nested" in i.title for i in report.items)


def test_validation_without_gcp_checks():
    report = SpecValidator(ClientFactory()).validate(ProcessSpec.model_validate(load_sample("sales_daily")), gcp_checks=False)
    assert report.ok and not report.gcp_checks


def test_lab_and_prod_sharing_a_write_location_is_flagged():
    spec_dict = load_sample("strategic_template")
    for v in spec_dict["variables"]:
        if v["name"] == "write_project":
            v["values"]["value"] = v["values"]["lab"]
    report = SpecValidator(ClientFactory(bq=FakeBq(), gcs=FakeGcs())).validate(ProcessSpec.model_validate(spec_dict), gcp_checks=False)
    assert any("Lab and prod share a write location: write_project" in i.title for i in report.items)


def test_copy_steps_are_checked_in_lab():
    spec = ProcessSpec.model_validate(load_sample("engagement_file_ingest"))
    gcs = FakeGcs({"demo_vendor_drop_test": {"outbound/Daily Engagement File_20260114.csv": CSV}, "demo_file_test_bucket": {}})
    report = _validator(gcs=gcs, bq=FakeBq(datasets={"demo-analytics-dev.clubwide"})).validate(spec)
    copy_items = [i for i in report.items if i.step == "move_vendor_files"]
    assert any(i.title.startswith("gcloud storage mv gs://demo_vendor_drop_test/outbound/") for i in copy_items)
    assert any("1 object(s) under gs://demo_vendor_drop_test/outbound/Daily Engagement File_" in i.title for i in copy_items)
    assert "value-gcp-vendor-drop" not in " ".join(gcs.calls)


def test_sql_logger_table_is_checked_in_the_adc_project():
    spec = ProcessSpec.model_validate(load_sample("strategic_template"))
    report = _validator(adc_project="my-gcloud-project").validate(spec)
    item = next(i for i in report.items if i.id == "sqllog")
    assert item.level == "warn" and "my-gcloud-project.pipeline_audit.pipeline_sql_log" in item.title
    ok = FakeBq(tables={"my-gcloud-project.pipeline_audit.pipeline_sql_log": SimpleNamespace()})
    assert next(i for i in _validator(bq=ok, adc_project="my-gcloud-project").validate(spec).items if i.id == "sqllog").level == "ok"


# ---------------------------------------------------------------- TXT (delimited text)
def test_txt_pipe_delimited_inference_detects_the_delimiter():
    r = infer_from_bytes(b"cust_id|name|joined\n7|Ann|2026-01-02\n8|Bo|2026-01-03\n", "TXT")
    assert [(c.name, c.type) for c in r.columns] == [("cust_id", "INTEGER"), ("name", "STRING"), ("joined", "DATE")]


# ---------------------------------------------------------------- Secret Manager
def test_secret_checks_read_metadata_only_in_the_lab_gsm_project():
    from fakes import FakeSecretManager

    spec_dict = load_sample("customer_snapshot")
    spec_dict["variables"].append({"name": "gsm_project", "group": "project", "values": {"lab": "gsm-dev", "value": "gsm-prod"}})
    spec_dict["secrets"] = [
        {"name": "ok_secret", "secret_id": "good"},
        {"name": "pinned_off", "secret_id": "rotated", "version": "2"},
        {"name": "missing", "secret_id": "nope"},
    ]
    spec = ProcessSpec.model_validate(spec_dict)
    sm = FakeSecretManager({
        "projects/gsm-dev/secrets/good": [(1, "ENABLED"), (2, "ENABLED")],
        "projects/gsm-dev/secrets/rotated": [(1, "ENABLED"), (2, "DISABLED")],
    })
    report = SpecValidator(ClientFactory(bq=FakeBq(), gcs=FakeGcs(), secrets=sm)).validate(spec)
    items = {i.id: i for i in report.items if i.id.startswith("secret:")}
    assert items["secret:ok_secret"].level == "ok" and "version 2 is enabled" in items["secret:ok_secret"].title
    assert items["secret:pinned_off"].level == "error"
    assert items["secret:missing"].level == "warn" and "secretAccessor" in items["secret:missing"].detail


# ---------------------------------------------------------------- email notification
def test_email_summary_names_the_adc_sender_not_the_relay():
    spec = ProcessSpec.model_validate(load_sample("strategic_template"))
    report = _validator(adc_email="pipeline.developer@example.com").validate(spec, gcp_checks=False)
    items = {i.id: i for i in report.items if i.id.startswith("email:")}
    summary = items["email:summary"]
    assert "from pipeline.developer@example.com" in summary.title
    assert "smtp.example.com" not in summary.title + summary.detail  # the relay is a hub setting, not shown
    assert "1 in lab, 2 in value" in summary.title and "email:sender" not in items
    no_sender = _validator().validate(spec, gcp_checks=False)
    assert any(i.id == "email:sender" and i.level == "warn" for i in no_sender.items)


FIXED_LINES = b"0000000001Ann Smith                     A 2026-01-15      125.50   12\n" \
              b"0000000002Bob Jones                     I 2025-12-01        0.00    3\n"


def test_fixed_width_layout_is_parsed_from_pasted_lines():
    cols = columns_from_layout("Field  Length  Type\nMember ID 10 STRING\nmember_name,30\nstatus_code\t2\n"
                               "join_date 10 date\nbalance 12 NUMERIC\nvisits 5 INT64\n")
    assert [(c.name, c.type, c.width) for c in cols] == [
        ("Member_ID", "STRING", 10), ("member_name", "STRING", 30), ("status_code", "STRING", 2),
        ("join_date", "DATE", 10), ("balance", "NUMERIC", 12), ("visits", "INTEGER", 5),
    ]
    with pytest.raises(ValueError, match="line 2"):
        columns_from_layout("a 5\nb\n")


def test_fixed_width_preview_splits_lines_and_flags_lengths():
    cols = load_sample("fixed_width_ingest")["steps"][0]["columns"]
    columns = [Column.model_validate(c) for c in cols]
    rows, problems = preview_fixed_width(FIXED_LINES, columns)
    assert problems == [] and rows[0]["member_name"] == "Ann Smith" and rows[1]["visits"] == "3"
    _rows, problems = preview_fixed_width(FIXED_LINES.replace(b"   12\n", b"\n"), columns)
    assert "shorter than the layout (69 characters" in problems[0]
    _rows, problems = preview_fixed_width(FIXED_LINES.replace(b"   12\n", b"   12 EXTRA\n"), columns)
    assert "longer than the layout" in problems[0]


def test_fixed_width_lab_sample_is_checked_against_the_layout():
    spec = ProcessSpec.model_validate(load_sample("fixed_width_ingest"))
    bq = FakeBq(datasets={"demo-analytics-dev.clubwide"})

    def report(data: bytes):
        gcs = FakeGcs({"demo_file_test_bucket": {"member_extract/MBR_EXTRACT_20260114.txt": data}})
        return SpecValidator(ClientFactory(bq=bq, gcs=gcs)).validate(spec)

    items = {i.id: i for i in report(FIXED_LINES).items}
    assert items["ingest_member_extract:layout"].level == "ok"
    bad = report(FIXED_LINES.replace(b"   12\n", b"\n"))
    layout = [i for i in bad.items if i.id.startswith("ingest_member_extract:layout:")]
    assert layout and layout[0].level == "warn" and "shorter than the layout" in layout[0].detail



# ---------------------------------------------------------------- Excel (.xlsx)
def _workbook(*rows, title: str = "Sheet1", extra_sheet: bool = False) -> bytes:
    import io

    from openpyxl import Workbook

    book = Workbook()
    book.active.title = title
    for row in rows:
        book.active.append(list(row))
    if extra_sheet:
        book.create_sheet("Notes")
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def test_excel_inference_reads_typed_cells_below_title_rows():
    data = _workbook(["Agent roster"], ["week 3"], ["Agent ID", "hired", "stamp", "active", "zip"],
                     [1001, dt.datetime(2026, 1, 15), dt.datetime(2026, 1, 15, 10, 30), True, "01234"],
                     title="Roster", extra_sheet=True)
    r = infer_from_bytes(data, "XLSX", sheet="Roster", skiprows=2, header_row=0)
    assert [(c.name, c.type) for c in r.columns] == [
        ("Agent_ID", "INTEGER"), ("hired", "DATE"), ("stamp", "DATETIME"), ("active", "BOOLEAN"), ("zip", "STRING")]
    assert r.detected["sheet"] == "Roster" and r.detected["sheets"] == ["Roster", "Notes"]
    with pytest.raises(ValueError, match="no sheet 'Other'"):
        infer_from_bytes(data, "XLSX", sheet="Other")
    with pytest.raises(ValueError, match="not a readable .xlsx"):
        infer_from_bytes(b"id,name\n1,x\n", "XLSX")


def _export(**kw) -> dict:
    step = {"type": "export", "function_name": "send_xlsx", "source_kind": "table", "source_table": "{write_project}.{ds}.t",
            "bucket": "out-dev", "file_name": "report_", "file_format": "XLSX"}
    step.update(kw)
    return step


def test_excel_export_is_one_file_without_compression():
    ProcessSpec.model_validate(load_sample("excel_ingest_export"))
    query = _export(source_kind="query", source_sql="SELECT 1 AS x", source_table=None)
    assert ProcessSpec.model_validate(minimal_spec(query)).steps[0].sharded is False  # EXPORT DATA shards; Excel does not
    for bad, message in (({"sharded": True}, "turn 'Sharded files' off"), ({"compression": "GZIP"}, "supports compression: NONE"),
                         ({"sheet_name": "a/b"}, "sheet name 'a/b' must be at most 31 characters")):
        with pytest.raises(ValueError, match=message):
            ProcessSpec.model_validate(minimal_spec(_export(**bad)))


def test_excel_fields_leave_older_spec_hashes_alone():
    spec = ProcessSpec.model_validate(minimal_spec(_export(file_format="CSV")))
    data = spec.model_dump(mode="json")
    assert data["steps"][0]["sheet_name"] == ""
    named = ProcessSpec.model_validate(minimal_spec(_export(sheet_name="Totals")))
    assert named.spec_hash() != ProcessSpec.model_validate(minimal_spec(_export())).spec_hash()


# ====================================================================== BigQuery's guess (lab, 1,000 rows)
def _cols(*names: str) -> list[Column]:
    return [Column(name=n, type="STRING", description=f"{n} column") for n in names]


def test_inference_keeps_the_data_rows_for_bigquerys_guess():
    rows = "".join(f"{i},x{i}\n" for i in range(1500))
    r = infer_from_bytes(("id,code\n" + rows).encode(), "CSV")
    assert len(r.rows) == 1000 and r.rows[0] == ["0", "x0"]


def test_bigquery_guess_types_by_position_and_deletes_its_table():
    bq = FakeBq(guess_types=["INTEGER", "FLOAT", "DATE", "INTEGER"])
    rows = [["007", "3.5", "2026-01-01", "1"], ["010", "4", "2026-01-02"], ["", "", "", ""]]
    r = guess_types_with_bigquery(bq, "proj-dev.tmp", _cols("zip", "amount", "day", "n"), rows)
    assert [(c.name, c.type) for c in r.columns] == [("zip", "STRING"), ("amount", "FLOAT"), ("day", "DATE"), ("n", "INTEGER")]
    assert r.columns[0].description == "zip column"  # only the type changes
    assert any("leading zeros" in w for w in r.warnings)
    table, csv_text, config = bq.loads[0]
    assert table.startswith("proj-dev.tmp._dahub_type_guess_") and bq.deleted == [table] and table not in bq.tables
    assert csv_text.splitlines() == ["007,3.5,2026-01-01,1", "010,4,2026-01-02,"]  # no header, blank rows dropped, padded
    assert config.autodetect is True and config.skip_leading_rows == 0
    assert r.detected == {"guessed_by": "bigquery", "rows_loaded": 2}


def test_bigquery_guess_demotes_columns_whose_rows_were_rejected():
    bq = FakeBq(guess_types=["INTEGER", "INTEGER"],
                bad_rows=["Could not parse 'n/a' as INT64 for field int64_field_1 (position 1) starting at location 9"])
    r = guess_types_with_bigquery(bq, "proj-dev.tmp", _cols("a", "b"), [["1", "2"], ["2", "n/a"]])
    assert [c.type for c in r.columns] == ["INTEGER", "STRING"]
    assert any("1 of 2 sample rows" in w for w in r.warnings)


def test_bigquery_guess_deletes_the_table_when_the_load_fails():
    bq = FakeBq(load_error="400 bad CSV")
    with pytest.raises(RuntimeError, match="bad CSV"):
        guess_types_with_bigquery(bq, "proj-dev.tmp", _cols("a"), [["1"]])
    assert bq.deleted == [bq.loads[0][0]]


def test_bigquery_guess_reads_excel_and_fixed_width_rows():
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["id", "joined"])
    book.active.append([1, dt.datetime(2026, 1, 5)])
    buf = io.BytesIO()
    book.save(buf)
    assert infer_from_bytes(buf.getvalue(), "XLSX").rows == [["1", "2026-01-05"]]
    columns = [Column(name="id", type="STRING", width=3), Column(name="name", type="STRING", width=5)]
    assert fixed_width_rows(b"001alice\n002bob  \n", columns) == [["001", "alice"], ["002", "bob"]]
    with pytest.raises(ValueError, match="length"):
        fixed_width_rows(b"x", [Column(name="id", type="STRING")])


def test_bigquery_guess_only_uses_a_lab_dataset():
    lab = {"write_project": "proj-dev", "temporary_dataset": "tmp_lab"}
    prod = {"write_project": "proj-prod", "temporary_dataset": "tmp"}
    assert lab_guess_target("{temporary_dataset}", lab, {"value": prod}, ["{temporary_dataset}"]) == ("proj-dev.tmp_lab", None)
    shared = {"write_project": "proj-dev", "temporary_dataset": "tmp_lab"}  # prod writes the same project.dataset
    target, problem = lab_guess_target("{temporary_dataset}", lab, {"value": shared}, ["{temporary_dataset}"])
    assert "value writes too" in problem
    assert lab_guess_target("other_ds", lab, {"value": shared}, ["{temporary_dataset}"]) == ("proj-dev.other_ds", None)
    assert lab_guess_target("{missing}", lab, {}, [])[1] == "enter a lab dataset for the temporary table"
    assert "write_project" in lab_guess_target("tmp", {}, {}, [])[1]

