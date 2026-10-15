"""Turn whatever the user has — a sample file, a JSON schema, a CREATE TABLE, an existing table — into columns."""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from typing import Any

import sqlglot
from sqlglot import exp

from ..models.rules import fill_variables
from ..models.schema import Column, columns_from_bq_json, sanitize_column_name

_INT_RE = re.compile(r"^[+-]?\d{1,18}$")
_FLOAT_RE = re.compile(r"^[+-]?(\d+\.\d*|\.\d+|\d+)([eE][+-]?\d+)?$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}(\.\d{1,9})?)?$")
_TIMESTAMP_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}(\.\d{1,9})?)?\s*(Z|UTC|[+-]\d{2}:?\d{2})$", re.I)
_BOOL_VALUES = {"true", "false"}
_LEADING_ZERO_RE = re.compile(r"^[+-]?0\d")
GUESS_ROWS = 1000  # sample rows BigQuery's guess loads
_NUMBER_TYPES = {"INTEGER", "FLOAT", "NUMERIC", "BIGNUMERIC"}
_BQ_POSITION_RE = re.compile(r"\(position (\d+)\)")
_DATASET_RE = re.compile(r"^[A-Za-z0-9_]{1,1024}$")
# declared type -> sample-inferred types that load fine into it
_COMPATIBLE = {
    "FLOAT": {"INTEGER"},
    "NUMERIC": {"INTEGER", "FLOAT"},
    "BIGNUMERIC": {"INTEGER", "FLOAT"},
    "TIMESTAMP": {"DATE", "DATETIME"},
    "DATETIME": {"DATE"},
    "JSON": {"RECORD", "INTEGER", "FLOAT", "BOOLEAN"},
}


@dataclass
class InferResult:
    columns: list[Column]
    sample_rows: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    detected: dict[str, Any] = field(default_factory=dict)  # e.g. delimiter, header
    rows: list[list[str]] = field(default_factory=list)  # the data rows read (text cells, up to 1000): BigQuery's guess


# ================================================================ sample files
def infer_from_bytes(
    data: bytes,
    file_format: str,
    *,
    filename: str = "",
    delimiter: str | None = None,
    has_header: bool = True,
    truncated: bool = False,
    sheet: str | None = None,
    skiprows: Any = None,
    header_row: int | None = None,
) -> InferResult:
    """Columns of a sample file. Excel only: ``sheet`` ('' / None = the first), and the pandas ``skiprows`` /
    ``header_row`` (the header's row once the skipped rows are gone) the step reads it with."""
    fmt = (file_format or "").upper()
    if fmt in ("CSV", "TXT"):  # TXT = delimited text; the delimiter is detected when not given
        return _infer_csv(data, delimiter, has_header, truncated)
    if fmt in ("JSON", "NEWLINE_DELIMITED_JSON", "NDJSON"):
        return _infer_ndjson(data, truncated)
    if fmt == "XLSX":
        return _infer_xlsx(data, has_header, sheet, skiprows=skiprows, header_row=header_row)
    if fmt == "PARQUET":
        return _infer_arrow(data, "parquet")
    if fmt == "ORC":
        return _infer_arrow(data, "orc")
    if fmt == "AVRO":
        return _infer_avro(data)
    raise ValueError(f"cannot infer a schema from {file_format or filename}")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _infer_csv(data: bytes, delimiter: str | None, has_header: bool, truncated: bool) -> InferResult:
    text = _decode(data)
    if truncated and "\n" in text:
        text = text[: text.rfind("\n")]  # drop the partial last line
    warnings: list[str] = []
    if not delimiter:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;|\t").delimiter
        except csv.Error:
            delimiter = ","
            warnings.append("could not detect the delimiter; assumed ','")
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _infer_rows(rows, has_header, warnings, {"delimiter": delimiter})


def _infer_rows(rows: list[list[str]], has_header: bool, warnings: list[str], detected: dict[str, Any]) -> InferResult:
    """Columns from rows of text cells (delimited text, or an Excel sheet's cells written as text)."""
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        raise ValueError("the sample file is empty")
    width = max(len(r) for r in rows[:1000])
    used: set[str] = set()
    if has_header:
        header, body = rows[0], rows[1:1001]
        names = [sanitize_column_name(h or f"column_{i + 1}", used) for i, h in enumerate(header + [""] * (width - len(header)))]
        renamed = [f"{h!r} → {n}" for h, n in zip(header, names) if h.strip() != n]
        if renamed:
            warnings.append("renamed headers to valid BigQuery names: " + ", ".join(renamed[:10]))
    else:
        body = rows[:1000]
        names = [f"column_{i + 1}" for i in range(width)]
    columns = []
    for idx, name in enumerate(names):
        values = [r[idx] for r in body if idx < len(r)]
        columns.append(Column(name=name, type=_infer_scalar_type(values)))
    if not body:
        warnings.append("the sample only has a header row; every column was typed STRING")
    sample = [dict(zip(names, r)) for r in body[:10]]
    return InferResult(columns, sample, warnings, {**detected, "header": has_header, "rows_sampled": len(body)}, body)


def _excel_text(value: Any) -> str:
    """An Excel cell as the text the type inference reads (dates without a time of day are dates)."""
    import datetime as dt

    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _infer_xlsx(data: bytes, has_header: bool, sheet: str | None = None, *, skiprows: Any = None,
                header_row: int | None = None) -> InferResult:
    """The first sheet (or ``sheet``) of an Excel workbook: up to 1000 rows below the header."""
    skipped = set(range(skiprows)) if isinstance(skiprows, int) else set(skiprows or [])
    above = (header_row or 0) + len(skipped)
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001  (zipfile / openpyxl errors)
        raise ValueError(f"not a readable .xlsx workbook: {exc}") from exc
    try:
        names = book.sheetnames
        if sheet and sheet not in names:
            raise ValueError(f"the workbook has no sheet {sheet!r} (it has {', '.join(names)})")
        ws = book[sheet] if sheet else book[names[0]]
        rows = [[_excel_text(v) for v in row]
                for _, row in zip(range(1001 + int(has_header) + above), ws.iter_rows(values_only=True))]
    finally:
        book.close()
    rows = [r for i, r in enumerate(rows) if i not in skipped][header_row or 0:]  # pandas: skiprows first, then header
    warnings = [f"read the sheet {ws.title!r} (the workbook has {len(names)} sheets)"] if len(names) > 1 else []
    return _infer_rows(rows, has_header, warnings, {"sheet": ws.title, "sheets": names})


def _infer_scalar_type(values: list[str]) -> str:
    vals = [v.strip() for v in values if v is not None and v.strip() != "" and v.strip().upper() not in ("NULL", "\\N")]
    if not vals:
        return "STRING"
    if any(_LEADING_ZERO_RE.match(v) for v in vals):
        return "STRING"  # codes like 01234 must keep their leading zeros
    checks = [
        ("BOOLEAN", lambda v: v.lower() in _BOOL_VALUES),
        ("INTEGER", lambda v: bool(_INT_RE.match(v))),
        ("FLOAT", lambda v: bool(_FLOAT_RE.match(v))),
        ("DATE", lambda v: bool(_DATE_RE.match(v))),
        ("TIMESTAMP", lambda v: bool(_TIMESTAMP_RE.match(v))),
        ("DATETIME", lambda v: bool(_DATETIME_RE.match(v))),
    ]
    for bq_type, check in checks:
        if all(check(v) for v in vals):
            return bq_type
    return "STRING"


def _infer_ndjson(data: bytes, truncated: bool) -> InferResult:
    text = _decode(data)
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if truncated and lines:
        lines = lines[:-1]
    records, warnings = [], []
    for i, line in enumerate(lines[:1000]):
        try:
            obj = json.loads(line)
        except json.JSONDecodeError:
            if i == 0 and text.lstrip().startswith("["):
                raise ValueError("this looks like a JSON array; BigQuery needs newline-delimited JSON (one object per line)") from None
            warnings.append(f"line {i + 1} is not valid JSON and was ignored")
            continue
        if isinstance(obj, dict):
            records.append(obj)
    if not records:
        raise ValueError("no JSON objects found in the sample")
    columns = _columns_from_records(records, set())
    return InferResult(columns, records[:10], warnings, {"rows_sampled": len(records)})


def _columns_from_records(records: list[dict[str, Any]], used: set[str]) -> list[Column]:
    keys: dict[str, list[Any]] = {}
    for rec in records:
        for k, v in rec.items():
            keys.setdefault(k, []).append(v)
    columns = []
    for key, values in keys.items():
        name = sanitize_column_name(key, used)
        columns.append(_column_from_values(name, values))
    return columns


def _column_from_values(name: str, values: list[Any]) -> Column:
    non_null = [v for v in values if v is not None]
    if non_null and all(isinstance(v, list) for v in non_null):
        items = [i for v in non_null for i in v if i is not None]
        inner = _column_from_values(name, items) if items else Column(name=name, type="STRING")
        return inner.model_copy(update={"mode": "REPEATED"})
    if non_null and all(isinstance(v, dict) for v in non_null):
        fields = _columns_from_records(non_null, set())
        return Column(name=name, type="RECORD", fields=fields or [Column(name="value", type="STRING")])
    if non_null and all(isinstance(v, bool) for v in non_null):
        return Column(name=name, type="BOOLEAN")
    if non_null and all(isinstance(v, int) and not isinstance(v, bool) for v in non_null):
        return Column(name=name, type="INTEGER")
    if non_null and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in non_null):
        return Column(name=name, type="FLOAT")
    if non_null and all(isinstance(v, str) for v in non_null):
        # JSON strings stay strings (zip codes, ids) unless they are clearly dates/timestamps
        inferred = _infer_scalar_type(non_null)
        return Column(name=name, type=inferred if inferred in ("DATE", "TIMESTAMP", "DATETIME") else "STRING")
    return Column(name=name, type="JSON" if any(isinstance(v, (dict, list)) for v in non_null) else "STRING")


def _arrow_to_bq(name: str, t: Any) -> Column:
    import pyarrow as pa

    if pa.types.is_struct(t):
        return Column(name=name, type="RECORD", fields=[_arrow_to_bq(f.name, f.type) for f in t])
    if pa.types.is_list(t) or pa.types.is_large_list(t):
        return _arrow_to_bq(name, t.value_type).model_copy(update={"mode": "REPEATED"})
    mapping = [
        (pa.types.is_boolean, "BOOLEAN"),
        (pa.types.is_integer, "INTEGER"),
        (pa.types.is_floating, "FLOAT"),
        (pa.types.is_decimal, "NUMERIC"),
        (pa.types.is_date, "DATE"),
        (pa.types.is_timestamp, "TIMESTAMP"),
        (pa.types.is_time, "TIME"),
        (pa.types.is_binary, "BYTES"),
        (pa.types.is_large_binary, "BYTES"),
    ]
    for check, bq in mapping:
        if check(t):
            if bq == "TIMESTAMP" and getattr(t, "tz", None) is None:
                bq = "DATETIME"
            if bq == "NUMERIC" and getattr(t, "precision", 0) > 38:
                bq = "BIGNUMERIC"
            return Column(name=name, type=bq)
    return Column(name=name, type="STRING")


def _infer_arrow(data: bytes, kind: str) -> InferResult:
    import pyarrow as pa

    buf = pa.BufferReader(data)
    try:
        if kind == "parquet":
            import pyarrow.parquet as pq

            schema = pq.read_schema(buf)
        else:
            import pyarrow.orc as orc

            schema = orc.ORCFile(buf).schema
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not read the {kind} file (a complete file is needed, not a partial sample): {exc}") from exc
    used: set[str] = set()
    columns = []
    for f in schema:
        col = _arrow_to_bq(sanitize_column_name(f.name, used), f.type)
        if not f.nullable and col.mode == "NULLABLE":
            col = col.model_copy(update={"mode": "REQUIRED"})
        columns.append(col)
    return InferResult(columns, [], [], {"format": kind})


# ---------------------------------------------------------------- avro (header only, no extra dependency)
def _read_long(buf: io.BytesIO) -> int:
    shift, result = 0, 0
    while True:
        b = buf.read(1)
        if not b:
            raise ValueError("truncated Avro header")
        byte = b[0]
        result |= (byte & 0x7F) << shift
        if not byte & 0x80:
            break
        shift += 7
    return (result >> 1) ^ -(result & 1)


def _infer_avro(data: bytes) -> InferResult:
    buf = io.BytesIO(data)
    if buf.read(4) != b"Obj\x01":
        raise ValueError("not an Avro object container file")
    meta: dict[str, bytes] = {}
    while True:
        count = _read_long(buf)
        if count == 0:
            break
        if count < 0:
            _read_long(buf)
            count = -count
        for _ in range(count):
            key = buf.read(_read_long(buf)).decode()
            meta[key] = buf.read(_read_long(buf))
    schema = json.loads(meta.get("avro.schema", b"{}"))
    if schema.get("type") != "record":
        raise ValueError("the Avro schema is not a record")
    used: set[str] = set()
    return InferResult([_avro_field(f, used) for f in schema.get("fields", [])], [], [], {"format": "avro"})


_AVRO_TYPES = {"string": "STRING", "bytes": "BYTES", "int": "INTEGER", "long": "INTEGER", "float": "FLOAT",
               "double": "FLOAT", "boolean": "BOOLEAN", "null": "STRING"}
_AVRO_LOGICAL = {"date": "DATE", "timestamp-millis": "TIMESTAMP", "timestamp-micros": "TIMESTAMP",
                 "local-timestamp-millis": "DATETIME", "local-timestamp-micros": "DATETIME",
                 "time-millis": "TIME", "time-micros": "TIME", "decimal": "NUMERIC", "uuid": "STRING"}


def _avro_field(f: dict[str, Any], used: set[str]) -> Column:
    name = sanitize_column_name(f["name"], used)
    t = f["type"]
    mode = "REQUIRED"
    if isinstance(t, list):
        non_null = [x for x in t if x != "null"]
        mode = "NULLABLE" if "null" in t else "REQUIRED"
        t = non_null[0] if len(non_null) == 1 else "string"
    col = _avro_type(name, t)
    return col.model_copy(update={"mode": "REPEATED" if col.mode == "REPEATED" else mode, "description": f.get("doc", "")})


def _avro_type(name: str, t: Any) -> Column:
    if isinstance(t, dict):
        if t.get("logicalType") in _AVRO_LOGICAL:
            return Column(name=name, type=_AVRO_LOGICAL[t["logicalType"]])
        if t.get("type") == "record":
            used: set[str] = set()
            return Column(name=name, type="RECORD", fields=[_avro_field(x, used) for x in t.get("fields", [])])
        if t.get("type") == "array":
            return _avro_type(name, t.get("items")).model_copy(update={"mode": "REPEATED"})
        if t.get("type") == "enum":
            return Column(name=name, type="STRING")
        return _avro_type(name, t.get("type"))
    return Column(name=name, type=_AVRO_TYPES.get(str(t), "STRING"))


# ================================================================ other schema sources
def columns_from_json_text(text: str) -> list[Column]:
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"not valid JSON: {exc}") from exc
    return columns_from_bq_json(data)


_SQL_TO_BQ = {
    "VARCHAR": "STRING", "CHAR": "STRING", "TEXT": "STRING", "NVARCHAR": "STRING", "STRING": "STRING",
    "INT": "INTEGER", "INT64": "INTEGER", "BIGINT": "INTEGER", "SMALLINT": "INTEGER", "TINYINT": "INTEGER",
    "INTEGER": "INTEGER", "FLOAT": "FLOAT", "DOUBLE": "FLOAT", "FLOAT64": "FLOAT", "REAL": "FLOAT",
    "DECIMAL": "NUMERIC", "NUMERIC": "NUMERIC", "BIGDECIMAL": "BIGNUMERIC", "BIGNUMERIC": "BIGNUMERIC",
    "BOOLEAN": "BOOLEAN", "BOOL": "BOOLEAN", "DATE": "DATE", "DATETIME": "DATETIME", "TIMESTAMP": "TIMESTAMP",
    "TIMESTAMPTZ": "TIMESTAMP", "TIME": "TIME", "BINARY": "BYTES", "VARBINARY": "BYTES", "BYTES": "BYTES",
    "JSON": "JSON", "GEOGRAPHY": "GEOGRAPHY",
}


def columns_from_ddl(ddl: str) -> list[Column]:
    try:
        create = sqlglot.parse_one(ddl, read="bigquery")
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not parse the CREATE TABLE statement: {str(exc).splitlines()[0]}") from exc
    schema = create.find(exp.Schema) if create else None
    defs = [c for c in (schema.expressions if schema else []) if isinstance(c, exp.ColumnDef)]
    if not defs:
        raise ValueError("no column definitions found; paste a CREATE TABLE ... (col TYPE, ...) statement")
    return [_column_def(c) for c in defs]


def _column_def(c: exp.ColumnDef) -> Column:
    mode = "REQUIRED" if any(isinstance(k.kind, exp.NotNullColumnConstraint) and not k.kind.args.get("allow_null")
                             for k in c.constraints) else "NULLABLE"
    desc = ""
    for k in c.constraints:
        if isinstance(k.kind, exp.Properties) or "description" in k.sql().lower():
            m = re.search(r"description\s*=\s*(['\"])(.*?)\1", k.sql(), re.I | re.S)
            if m:
                desc = m.group(2)
    col = _data_type(c.name, c.args.get("kind"))
    return col.model_copy(update={"mode": "REPEATED" if col.mode == "REPEATED" else mode, "description": desc})


def _data_type(name: str, dtype: exp.DataType | None) -> Column:
    if dtype is None:
        return Column(name=name, type="STRING")
    t = dtype.this
    if t == exp.DataType.Type.ARRAY:
        inner = dtype.expressions[0] if dtype.expressions else None
        return _data_type(name, inner).model_copy(update={"mode": "REPEATED"})
    if t == exp.DataType.Type.STRUCT:
        fields = []
        for e in dtype.expressions:
            if isinstance(e, exp.ColumnDef):
                fields.append(_column_def(e))
            elif isinstance(e, exp.DataType):
                fields.append(_data_type(f"f{len(fields)}", e))
        return Column(name=name, type="RECORD", fields=fields or [Column(name="value", type="STRING")])
    sql_name = dtype.sql(dialect="bigquery").split("(")[0].split("<")[0].upper()
    return Column(name=name, type=_SQL_TO_BQ.get(sql_name, _SQL_TO_BQ.get(t.name, "STRING")))


def columns_from_bq_table(bq_client: Any, table_id: str) -> list[Column]:
    table = bq_client.get_table(table_id)
    return columns_from_bq_json([f.to_api_repr() for f in table.schema])


# ================================================================ BigQuery's own guess (lab)
def lab_guess_target(dataset: str, lab: dict[str, str], others: dict[str, dict[str, str]],
                     prod_datasets: list[str]) -> tuple[str, str | None]:
    """Where BigQuery's guess may create its temporary table: ``<lab write_project>.<dataset>`` (``dataset`` may be a
    {variable}). Returns (project.dataset, problem). Refused when another environment (prod) writes to that dataset:
    ``others`` = {env: variable values}, ``prod_datasets`` = the step's dataset templates."""
    project = (lab.get("write_project") or "").strip()
    resolved = fill_variables(dataset or "", lab).strip()
    if not project:
        return "", "enter the lab value of write_project first (Variables)"
    if not resolved or "{" in resolved:
        return "", "enter a lab dataset for the temporary table"
    if not _DATASET_RE.match(resolved):
        return "", f"{resolved!r} is not a dataset name"
    target = f"{project}.{resolved}"
    for env, values in others.items():
        theirs = {f"{values.get('write_project', '')}.{fill_variables(t, values)}" for t in prod_datasets if t}
        if target in theirs:
            return target, f"{target} is where {env} writes too: pick a dataset only lab uses"
    return target, None


def guess_types_with_bigquery(bq_client: Any, table_prefix: str, columns: list[Column], rows: list[list[str]], *,
                              location: str | None = None) -> InferResult:
    """BigQuery's own guess of the column types: load up to 1,000 sample rows (CSV, autodetect) into a temporary table
    in a lab dataset (``project.dataset``), read the types it detected, then delete the table. Each column keeps its
    name, length and description; only the type changes."""
    import uuid

    from google.cloud import bigquery

    rows = [r for r in rows if any(str(c or "").strip() for c in r)][:GUESS_ROWS]
    if not columns:
        raise ValueError("read the columns first")
    if not rows:
        raise ValueError("the sample has no data rows for BigQuery to guess from")
    width = len(columns)
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for r in rows:
        writer.writerow((list(r) + [""] * width)[:width])
    table_id = f"{table_prefix}._dahub_type_guess_{uuid.uuid4().hex[:8]}"
    config = bigquery.LoadJobConfig(
        source_format=bigquery.SourceFormat.CSV, autodetect=True, skip_leading_rows=0,  # 0 = no header row
        allow_quoted_newlines=True, max_bad_records=len(rows), write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
    )
    warnings: list[str] = []
    try:
        job = bq_client.load_table_from_file(io.BytesIO(buf.getvalue().encode("utf-8")), table_id, job_config=config,
                                             location=location)
        job.result(timeout=120)
        detected = columns_from_bq_json([f.to_api_repr() for f in bq_client.get_table(table_id).schema])
    finally:
        try:
            bq_client.delete_table(table_id, not_found_ok=True)
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"could not delete the temporary table {table_id} ({exc}): delete it by hand")
    if len(detected) != width:
        raise ValueError(f"BigQuery found {len(detected)} columns but the schema has {width}")
    demoted: set[int] = set()
    loaded = getattr(job, "output_rows", None)
    if loaded is not None and loaded < len(rows):
        errors = [str(e.get("message", "")) for e in getattr(job, "errors", None) or []]
        demoted = {int(m.group(1)) for m in map(_BQ_POSITION_RE.search, errors) if m}
        warnings.append(f"{len(rows) - loaded} of {len(rows)} sample rows did not fit BigQuery's guess"
                        + (f" (e.g. {errors[0][:200]})" if errors else ""))
    out = []
    for i, (col, guess) in enumerate(zip(columns, detected)):
        new_type = guess.type
        if i in demoted and new_type != "STRING":
            warnings.append(f"{col.name}: some values are not {new_type}; kept STRING")
            new_type = "STRING"
        elif new_type in _NUMBER_TYPES and any(_LEADING_ZERO_RE.match(str(r[i]).strip()) for r in rows if i < len(r)):
            warnings.append(f"{col.name}: BigQuery guessed {new_type}; kept STRING to keep the leading zeros")
            new_type = "STRING"
        out.append(col.model_copy(update={"type": new_type}))
    return InferResult(out, [], warnings, {"guessed_by": "bigquery", "rows_loaded": len(rows)})


# ================================================================ sample vs schema
def compare_with_sample(columns: list[Column], sample: InferResult, file_format: str) -> list[str]:
    """Human-readable mismatches between the declared schema and a sample file."""
    problems = []
    declared = [c.name.lower() for c in columns]
    found = [c.name.lower() for c in sample.columns]
    if file_format.upper() in ("CSV", "TXT", "XLSX"):  # read by position (names come from the schema)
        if len(declared) != len(found):
            problems.append(f"the schema has {len(declared)} columns but the sample file has {len(found)}")
        elif sample.detected.get("header"):
            diffs = [f"#{i + 1}: schema '{d}' vs file '{f}'" for i, (d, f) in enumerate(zip(declared, found)) if d != f]
            if diffs:
                problems.append("column order/names differ from the file header (CSV loads by position): " + "; ".join(diffs[:8]))
    else:
        missing = [d for d in declared if d not in found]
        extra = [f for f in found if f not in declared]
        if missing:
            problems.append(f"in the schema but not in the sample file: {', '.join(missing[:10])}")
        if extra:
            problems.append(f"in the sample file but not in the schema (needs ignore_unknown_values): {', '.join(extra[:10])}")
    by_name = {c.name.lower(): c for c in sample.columns}
    for c in columns:
        s = by_name.get(c.name.lower())
        if not s or s.type == c.type or c.type in ("STRING", "BYTES", "GEOGRAPHY", "JSON"):
            continue
        if s.type == "STRING":
            bad = [str(r.get(s.name)) for r in sample.sample_rows if r.get(s.name) not in (None, "")][:3]
            if bad:
                problems.append(f"column {c.name}: schema says {c.type}, but sample values look like text (e.g. {', '.join(bad)})")
        elif s.type not in _COMPATIBLE.get(c.type, set()):
            problems.append(f"column {c.name}: schema says {c.type}, sample values look like {s.type}")
    return problems



# ================================================================ fixed-width layouts
_LAYOUT_SPLIT_RE = re.compile(r"[,\t;|]+|\s+")


def columns_from_layout(text: str) -> list[Column]:
    """A fixed-width layout pasted one column per line: ``name length [type]`` (spaces, tabs or commas between).
    A first line such as ``name length type`` is taken as a header and skipped."""
    columns: list[Column] = []
    problems: list[str] = []
    used: set[str] = set()
    for n, raw in enumerate((text or "").splitlines(), 1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        parts = [p for p in _LAYOUT_SPLIT_RE.split(line) if p]
        k = next((i for i, p in enumerate(parts) if p.isdigit()), None)  # the name may contain spaces: "Member ID 10"
        if not k:
            if k is None and not columns and not problems:
                continue  # a header line such as "name length type"
            problems.append(f"line {n}: expected 'name length [type]', got {line!r}")
            continue
        name, width = " ".join(parts[:k]), int(parts[k])
        if width < 1:
            problems.append(f"line {n}: the length of {name} must be at least 1")
            continue
        try:
            columns.append(Column(name=sanitize_column_name(name, used), type=parts[k + 1] if len(parts) > k + 1 else "STRING",
                                  width=width))
        except ValueError as exc:
            problems.append(f"line {n}: {str(exc).splitlines()[-1]}")
    if problems:
        raise ValueError("; ".join(problems[:5]) + (f" (and {len(problems) - 5} more)" if len(problems) > 5 else ""))
    if not columns:
        raise ValueError("no columns found: write one column per line, e.g. member_id 10 STRING")
    return columns


def pandas_rows(value: str) -> Any:
    """A pandas header / skiprows setting: None, a row number or a list of rows."""
    value = (value or "None").strip()
    return None if value.lower() == "none" else json.loads(value)


def _read_fixed_width(data: bytes, columns: list[Column], header: str, skiprows: str, rows: int,
                      truncated: bool) -> tuple[str, Any, Any, Any]:
    """(text, header row, skiprows, the first ``rows`` lines split by the column lengths, as the script reads them)."""
    import pandas as pd

    text = _decode(data)
    if truncated and "\n" in text:
        text = text[: text.rfind("\n")]  # drop the partial last line
    header_row, skip = pandas_rows(header), pandas_rows(skiprows)
    df = pd.read_fwf(io.StringIO(text), widths=[c.width or 0 for c in columns], names=[c.name for c in columns],
                     header=header_row, skiprows=skip, nrows=rows, dtype=str, keep_default_na=False)
    return text, header_row, skip, df


def fixed_width_rows(data: bytes, columns: list[Column], *, header: str = "None", skiprows: str = "None",
                     truncated: bool = False) -> list[list[str]]:
    """The first 1,000 data lines of a fixed-width sample split by the column lengths (for BigQuery's guess)."""
    if not columns or not all(c.width for c in columns):
        raise ValueError("enter the length of every column first")
    _text, _header, _skip, df = _read_fixed_width(data, columns, header, skiprows, GUESS_ROWS, truncated)
    return [[str(v) for v in rec] for rec in df.itertuples(index=False, name=None)]


def preview_fixed_width(data: bytes, columns: list[Column], *, header: str = "None", skiprows: str = "None",
                        rows: int = 5, truncated: bool = False) -> tuple[list[dict[str, str]], list[str]]:
    """Split a sample fixed-width file by the column widths, as the generated script will.
    Returns (the first rows, problems such as lines whose length differs from the layout)."""
    widths = [c.width or 0 for c in columns]
    if not columns or not all(widths):
        return [], ["enter the length of every column first"]
    text, header_row, skip, df = _read_fixed_width(data, columns, header, skiprows, rows, truncated)
    out = [{k: str(v) for k, v in rec.items()} for rec in df.to_dict(orient="records")]

    lines = text.splitlines()
    skipped = set(range(skip)) if isinstance(skip, int) else set(skip or [])
    kept = [ln for i, ln in enumerate(lines) if i not in skipped]
    if header_row is not None:
        kept = kept[int(header_row) + 1:]
    data_lines = [ln for ln in kept[:1000] if ln.strip()]
    total = sum(widths)
    problems: list[str] = []
    if not data_lines:
        problems.append("the sample has no data lines after the header and skipped rows")
    longer = [len(ln) for ln in data_lines if len(ln.rstrip()) > total]
    shorter = [len(ln) for ln in data_lines if len(ln) < total]
    if longer:
        problems.append(f"{len(longer)} of {len(data_lines)} lines are longer than the layout ({total} characters; "
                        f"up to {max(longer)}): a column may be missing or a length too small")
    if shorter:
        problems.append(f"{len(shorter)} of {len(data_lines)} lines are shorter than the layout ({total} characters; "
                        f"as short as {min(shorter)}): the last columns are read as empty")
    return out, problems
